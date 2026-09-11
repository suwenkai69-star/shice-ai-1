from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from data_foundation import (
    PAYMENT_ALIASES,
    SALES_ALIASES,
    SETTLEMENT_ALIASES,
    _decode_text,
)


class FilePreviewError(ValueError):
    def __init__(self, message: str, code: str = "FILE_PARSE_ERROR"):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class FilePreview:
    data_type: str
    row_count: int
    headers: tuple[str, ...]
    rows: tuple[dict[str, Any], ...]
    detected_channel: str | None
    report_scope: str
    business_date: str | None

    def summary(self) -> dict[str, Any]:
        return {
            "data_type": self.data_type,
            "row_count": self.row_count,
            "headers": list(self.headers),
            "detected_channel": self.detected_channel,
            "report_scope": self.report_scope,
            "business_date": self.business_date,
        }


def _serializable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None:
        return ""
    return value


def _xlsx_rows(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise FilePreviewError("XLSX 文件无法读取，请重新从平台导出后再上传") from exc
    try:
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            return [], []
        headers = [str(x).strip() if x is not None else "" for x in header_row]
        if not any(headers):
            return [], []
        rows: list[dict[str, Any]] = []
        for values in iterator:
            row = {headers[i]: _serializable(values[i]) if i < len(values) else "" for i in range(len(headers)) if headers[i]}
            if any(str(v).strip() for v in row.values() if v is not None):
                rows.append(row)
        return rows, [h for h in headers if h]
    finally:
        wb.close()


def _csv_rows(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        text = _decode_text(data)
    except Exception as exc:
        raise FilePreviewError(str(exc)) from exc
    reader = csv.DictReader(io.StringIO(text))
    rows = [dict(row) for row in reader]
    return rows, list(reader.fieldnames or [])


def _has_any(headers: set[str], aliases: list[str]) -> bool:
    return bool(headers.intersection(aliases))


def _pick_alias(row: dict[str, Any], aliases: list[str]) -> Any:
    for alias in aliases:
        if alias in row and str(row[alias]).strip() != "":
            return row[alias]
    return None


def _detect_type(headers: list[str]) -> str:
    h = {str(x).strip() for x in headers if x}
    if _has_any(h, PAYMENT_ALIASES["payment_date"]) and _has_any(h, PAYMENT_ALIASES["amount"]):
        return "PAYMENT"
    settlement_marker = (
        _has_any(h, SETTLEMENT_ALIASES["settlement_date"])
        or _has_any(h, SETTLEMENT_ALIASES["settlement_period_start"])
        or _has_any(h, SETTLEMENT_ALIASES["settlement_period_end"])
        or _has_any(h, SETTLEMENT_ALIASES["settlement_reference"])
    )
    if settlement_marker:
        return "SETTLEMENT"
    if _has_any(h, SALES_ALIASES["business_date"]) and _has_any(h, SALES_ALIASES["gross_sales"]):
        return "SALES"
    raise FilePreviewError("暂时无法判断这份文件属于销售、结算还是到账数据，请检查表头")


def preview_file(filename: str, data: bytes) -> FilePreview:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".xls":
        raise FilePreviewError("暂不支持旧版 XLS，请另存为 XLSX 或 CSV 后上传", code="UNSUPPORTED_XLS")
    if suffix == ".xlsx":
        rows, headers = _xlsx_rows(data)
    elif suffix == ".csv":
        rows, headers = _csv_rows(data)
    else:
        raise FilePreviewError("V1 文件上传仅支持 CSV 和 XLSX", code="UNSUPPORTED_FILE_TYPE")
    if not rows:
        raise FilePreviewError("文件里没有识别到可导入的数据")
    data_type = _detect_type(headers)

    aliases = SALES_ALIASES if data_type == "SALES" else PAYMENT_ALIASES if data_type == "PAYMENT" else SETTLEMENT_ALIASES
    channel_values = {
        str(v).strip()
        for row in rows
        if (v := _pick_alias(row, aliases["channel_code"])) is not None and str(v).strip()
    }
    channel = next(iter(channel_values)) if len(channel_values) == 1 else None

    report_scope = "CHANNEL"
    for row in rows:
        raw_scope = row.get("report_scope") or row.get("scope")
        if raw_scope and str(raw_scope).strip().upper() == "WHOLE_STORE":
            report_scope = "WHOLE_STORE"
            break
    day_aliases = SALES_ALIASES["business_date"] if data_type == "SALES" else PAYMENT_ALIASES["payment_date"] if data_type == "PAYMENT" else SETTLEMENT_ALIASES["settlement_date"]
    dates = {
        str(v).strip().replace("/", "-")
        for row in rows
        if (v := _pick_alias(row, day_aliases)) is not None and str(v).strip()
    }
    one_date = next(iter(dates)) if len(dates) == 1 else None
    return FilePreview(
        data_type=data_type,
        row_count=len(rows),
        headers=tuple(headers),
        rows=tuple(rows),
        detected_channel=channel,
        report_scope=report_scope,
        business_date=one_date,
    )
